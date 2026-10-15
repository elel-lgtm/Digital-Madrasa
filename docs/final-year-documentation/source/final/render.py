"""Renders final.json into a Word document in the academic style of Document B:
Times New Roman 12, 1.5 line spacing, justified; centred bold chapter titles on a
new page; bold centred captions below figures and tables; Roman page numbers on
the preliminary pages and Arabic from Chapter One; Word fields for the contents,
figure and table lists (updated by Word when the PDF is made).

Usage: python render.py final.json out.docx
"""
import json, re, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

SRC, OUT = sys.argv[1], sys.argv[2]
blocks = json.load(open(SRC, encoding='utf-8'))
FONT = 'Times New Roman'
TEXT_W = 6.7          # 8.5 in page, 1.0 in left and 0.8 in right margin (as in Document B)

doc = Document()
st = doc.styles


def set_font(style, size=12, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(k), FONT)
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if fonts.get(qn(k)) is not None:
            del fonts.attrib[qn(k)]


normal = st['Normal']
set_font(normal, 12)
pf = normal.paragraph_format
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.line_spacing = 1.5
pf.space_after = Pt(6)
pf.space_before = Pt(0)

for name, size, align, before, after in (('Heading 1', 14, WD_ALIGN_PARAGRAPH.CENTER, 0, 12),
                                         ('Heading 2', 12, WD_ALIGN_PARAGRAPH.LEFT, 12, 6),
                                         ('Heading 3', 12, WD_ALIGN_PARAGRAPH.LEFT, 10, 4)):
    s = st[name]
    set_font(s, size, bold=True, italic=False)
    s.paragraph_format.alignment = align
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.line_spacing = 1.5
    s.paragraph_format.keep_with_next = True
st['Heading 1'].paragraph_format.page_break_before = True

from docx.enum.style import WD_STYLE_TYPE
sub = st.add_style('Subheading', WD_STYLE_TYPE.PARAGRAPH)
sub.base_style = normal
set_font(sub, 12, bold=True, italic=True)
sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
sub.paragraph_format.keep_with_next = True
sub.paragraph_format.space_before = Pt(8)
for name in ('Figure Caption', 'Table Caption'):
    c = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    c.base_style = normal
    set_font(c, 12, bold=True)
    c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(4)
    c.paragraph_format.space_after = Pt(12)
    c.paragraph_format.line_spacing = 1.15
for name in ('TOC 1', 'TOC 2', 'TOC 3', 'Table of Figures'):
    try:
        s = st[name]
    except KeyError:
        s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
    s.element.set(qn('w:styleId'), name.replace(' ', '').replace('Tableof', 'Tableof'))
    s.element.find(qn('w:name')).set(qn('w:val'), name.lower() if name.startswith('TOC') else 'table of figures')
    set_font(s, 11, bold=(name == 'TOC 1'))
    s.paragraph_format.space_after = Pt(2)
    s.paragraph_format.line_spacing = 1.15
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.left_indent = Inches({'TOC 1': 0, 'TOC 2': 0.25, 'TOC 3': 0.5, 'Table of Figures': 0}[name])
set_font(st['List Bullet'], 12)
st['List Bullet'].paragraph_format.line_spacing = 1.5
st['List Bullet'].paragraph_format.space_after = Pt(3)

sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(1)
sec.left_margin, sec.right_margin = Inches(1), Inches(0.8)


# ------------------------------------------------------------------ helpers
def add_runs(p, runs, size=None):
    for text, bold, italic in runs:
        r = p.add_run(text)
        r.bold = bool(bold) or None
        r.italic = bool(italic) or None
        if size:
            r.font.size = Pt(size)
    return p


def field(p, code, placeholder=''):
    r = p.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    r = p.add_run()
    t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = f' {code} '; r._r.append(t)
    r = p.add_run()
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r._r.append(s)
    p.add_run(placeholder)
    r = p.add_run()
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)


def page_numbers(section, fmt):
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(footer.runs):
        r._r.getparent().remove(r._r)
    if fmt:
        field(footer, 'PAGE', '1')
        for r in footer.runs:
            r.font.size = Pt(11)
    pg = section._sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType'); section._sectPr.append(pg)
    if fmt:
        pg.set(qn('w:fmt'), fmt)
        pg.set(qn('w:start'), '1')


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '000000')
        b.append(e)
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
    c = OxmlElement('w:cantSplit'); c.set(qn('w:val'), 'true'); trPr.append(c)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement('w:cantSplit'); c.set(qn('w:val'), 'true'); trPr.append(c)


def table(rows, widths, font=10, header=True, bold_last=False):
    ncol = max(len(r) for r in rows)
    rows = [r + [''] * (ncol - len(r)) for r in rows]
    if not widths:
        longest = [max(len(r[c]) for r in rows) for c in range(ncol)]
        tot = sum(max(l, 6) for l in longest)
        widths = [TEXT_W * max(l, 6) / tot for l in longest]
    scale = min(1.0, (TEXT_W - 0.05) / sum(widths))
    widths = [w * scale for w in widths]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, r in enumerate(rows):
        row = t.rows[i]
        if i == 0 and header:
            repeat_header(row)
        else:
            cant_split(row)
        for j, text in enumerate(r):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if len(rows) <= 9:
                p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.size = Pt(font)
            if (i == 0 and header) or (bold_last and i == len(rows) - 1):
                run.bold = True
            if i == 0 and header:
                shade(cell, 'D9E2F3')
    for j, w in enumerate(widths):
        for cell in t.columns[j].cells:
            cell.width = Inches(w)
    return t


def caption(kind, number, title):
    p = doc.add_paragraph(style=f'{kind} Caption')
    p.add_run(f'{kind} {number} {title}')
    return p


fig_seen, tab_seen = [], []


def check_number(store, number):
    ch, n = number.split('.')
    prev = [x for x in store if x.split('.')[0] == ch]
    expected = len(prev) + 1
    assert int(n) == expected, f'{number} out of order (expected {ch}.{expected})'
    store.append(number)


# ------------------------------------------------------------------ render
for b in blocks:
    t = b['t']
    if t == 'cover':
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture('img_B/B001.png', width=Inches(2.6))
        for text, size, bold, before in (
                ('Digital Madrasa: An Integrated Application That Will Manage Madrasas in Kenya', 16, True, 30),
                ('MOHAMMEDNUR SUKYAN', 12, True, 30),
                ('BSCS/2024/52324', 12, True, 4),
                ('This project is submitted in partial fulfillment of the requirements for the award of Bachelor of Science in Computer Science.', 12, True, 30),
                ('School of Business and Technology', 12, True, 10),
                ('Umma University', 12, True, 0),
                ('Kajiado, Kenya', 12, True, 0),
                ('November 2026', 12, True, 30),
                ('This project report is available for library use on the understanding that it is copyright material and that no quotation from it may be published without proper acknowledgment.', 11, True, 110)):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(before)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.4)
            r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        page_numbers(doc.sections[0], None)
    elif t == 'section':
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        s.left_margin, s.right_margin = Inches(1), Inches(0.8)
        page_numbers(s, 'lowerRoman' if b['numbering'] == 'roman' else 'decimal')
    elif t == 'h1':
        p = doc.add_paragraph(b['text'], style='Heading 1')
        if len(doc.paragraphs) >= 2 and doc.paragraphs[-2]._p.find('.//' + qn('w:sectPr')) is not None:
            p.paragraph_format.page_break_before = False
    elif t == 'h2':
        doc.add_paragraph(b['text'], style='Heading 2')
    elif t == 'h3':
        doc.add_paragraph(b['text'], style='Heading 3')
    elif t == 'h4':
        doc.add_paragraph(b['text'], style='Subheading')
    elif t == 'p':
        p = doc.add_paragraph()
        add_runs(p, b['runs'])
        if b.get('keep_next'):
            p.paragraph_format.keep_with_next = True
        if 'indent' in b:
            p.paragraph_format.left_indent = Inches(b['indent'])
        if 'after' in b:
            p.paragraph_format.space_after = Pt(b['after'])
        if 'line' in b:
            p.paragraph_format.line_spacing = b['line']
    elif t == 'bullet':
        p = doc.add_paragraph(style='List Bullet')
        add_runs(p, b['runs'])
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif t == 'num':
        lst = b['list']
        globals().setdefault('_counters', {})
        _counters[lst] = _counters.get(lst, 0) + 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f'{_counters[lst]}.\t')
        add_runs(p, b['runs'])
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
    elif t == 'fig':
        check_number(fig_seen, b['number'])
        path = b['path']
        w, h = Image.open(path).size
        width = min(b['width'], TEXT_W - 0.1)
        if h / w * width > 7.6:      # keep a figure and its caption on one page
            width = 7.6 * w / h
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(path, width=Inches(width))
        caption('Figure', b['number'], b['title'])
    elif t == 'table':
        check_number(tab_seen, b['number'])
        table(b['rows'], b['widths'], font=b.get('font', 10),
              bold_last=b['rows'][-1][0].lower().startswith(('total', 'database security')) and b['number'] in ('A.3',))
        caption('Table', b['number'], b['title'])
    elif t == 'sign':
        def line(text, bold=False, after=10):
            p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(after)
            r = p.add_run(text); r.bold = bold or None
        dots = '.' * 40
        if b['role'] == 'student':
            line('Student’s signature:', True)
            line(f'Name {dots}   Signature {dots}')
            line(f'Date {dots}', after=24)
        else:
            line('Supervisor’s signature:', True)
            line(f'Signature {dots}   Date {dots}', after=18)
            line('MR. PATRICK MUCHIRI ALIANDA', after=0)
            line('School of Business and Technology, Umma University,', after=0)
            line('Kenya', after=0)
    elif t in ('toc', 'lof', 'lot'):
        title = {'toc': 'Table of Contents', 'lof': 'Table of Figures', 'lot': 'Table of Tables'}[t]
        p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title); r.bold = True; r.font.size = Pt(12)
        code = {'toc': 'TOC \\o "1-3" \\h \\z \\u', 'lof': 'TOC \\h \\z \\t "Figure Caption,4"', 'lot': 'TOC \\h \\z \\t "Table Caption,4"'}[t]
        p = doc.add_paragraph()
        field(p, code, 'Right-click and choose Update Field.')
    elif t == 'abbr':
        table([['Abbreviation', 'Meaning']] + [list(r) for r in b['rows']], [1.5, 4.6], font=11)
    elif t == 'refs':
        for item in b['items']:
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(8)
            p.add_run(item)
    else:
        raise ValueError(t)

core = doc.core_properties
core.title = 'Digital Madrasa: An Integrated Application That Will Manage Madrasas in Kenya'
core.author = 'Mohammednur Sukyan'
core.last_modified_by = 'Mohammednur Sukyan'
core.subject = 'Final-year project report, Bachelor of Science in Computer Science, Umma University'
core.keywords = 'Digital Madrasa; ILMORA; madrasa management; Quran memorisation; Kenya'
doc.save(OUT)
print('saved', OUT, '| figures', len(fig_seen), '| tables', len(tab_seen))
