"""Automated checks on the final report (docx after Word updated its fields, and the PDF)."""
import json, re, sys
import docx, pymupdf

DOCX, PDF = sys.argv[1], sys.argv[2]
d = docx.Document(DOCX)
paras = [p.text for p in d.paragraphs]
cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
full = '\n'.join(paras + cells)
results = []


def check(name, ok, detail=''):
    results.append((name, bool(ok), detail))
    print(('PASS ' if ok else 'FAIL ') + name + (f' — {detail}' if detail else ''))


# ---- placeholders and personal details
ph = sorted(set(re.findall(r'\[(?:INSERT|STUDENT|REGISTRATION|SUPERVISOR|MONTH|CONFIRM)[^\]]*\]', full)))
check('no placeholders remain', not ph, ', '.join(ph))
for s in ('Mohammednur Sukyan', 'MOHAMMEDNUR SUKYAN', 'BSCS/2024/52324', 'PATRICK MUCHIRI ALIANDA', 'Mr. Patrick Muchiri Alianda',
          'School of Business and Technology', 'Umma University', 'Kajiado, Kenya', 'November 2026'):
    check(f'contains "{s}"', s in full)
check('registration number not used as an app admission number', 'DM-2024-52324' not in full)

# ---- chapter four: exactly 12 analysed questions, 4 per group
body = '\n'.join(paras)
start = body.rindex('\n4.5.2 Questionnaire Analysis\n')      # the section itself, not its contents entry
end = body.rindex('\n4.6 Secondary Data Collection\n')
c4 = body[start:end]
groups = re.split(r'\n(?:Madrasa Owner and Teacher|Parent|Student) Questionnaire Analysis\n', c4)[1:]
counts = [len(re.findall(r'^Question \d+:', g, re.M)) for g in groups]
check('Chapter Four analyses 12 questions (4 owners/teachers, 4 parents, 4 students)', counts == [4, 4, 4], str(counts))
figs4 = [int(x) for x in re.findall(r'^Figure 4\.(\d+) ', c4, re.M)]
check('one chart per analysed question (Figures 4.1-4.12)', figs4 == list(range(1, 13)), str(figs4))
check('every analysed question has an Analysis and a System solution',
      c4.count('Analysis: ') == 12 and c4.count('System solution: ') == 12,
      f"{c4.count('Analysis: ')} / {c4.count('System solution: ')}")

# ---- figures and tables
caps_f = [p.text for p in d.paragraphs if p.style.name == 'Figure Caption']
caps_t = [p.text for p in d.paragraphs if p.style.name == 'Table Caption']
fig_nums = [re.match(r'Figure (\S+)', c).group(1) for c in caps_f]
tab_nums = [re.match(r'Table (\S+)', c).group(1) for c in caps_t]
prose = '\n'.join(p.text for p in d.paragraphs if p.style.name not in ('Figure Caption', 'Table Caption') and not p.style.name.startswith(('toc', 'TOC')))
refd_f = set(re.findall(r'Figures? ((?:\d|A)\.\d+)', prose))
refd_t = set(re.findall(r'Tables? ((?:\d|A)\.\d+)', prose))
check('figure captions', len(fig_nums) == len(set(fig_nums)), f'{len(fig_nums)} figures, numbers unique')
check('table captions', len(tab_nums) == len(set(tab_nums)), f'{len(tab_nums)} tables, numbers unique')
check('every figure mentioned in the text exists', refd_f <= set(fig_nums), str(sorted(refd_f - set(fig_nums))))
check('every table mentioned in the text exists', refd_t <= set(tab_nums), str(sorted(refd_t - set(tab_nums))))
imgs = len(d.inline_shapes)
check('every figure caption has an embedded image', imgs == len(fig_nums) + 1, f'{imgs} images (incl. the cover logo)')

# ---- citations against the reference list
refs_start = max(i for i, p in enumerate(paras) if p == 'REFERENCES')
refs = [p for p in paras[refs_start + 1:] if p.strip()]
text_part = '\n'.join(paras[:refs_start] + cells)


def key_of(ref):
    m = re.match(r'^(.+?)\s\((\d{4}|n\.d\.)', ref)
    author = m.group(1).strip().rstrip('.')
    # Personal authors start "Surname, I."; corporate authors are the whole name.
    personal = re.match(r"^[A-Z][\w\-’'. ]+, [A-Z]\.", m.group(1).strip())
    surname = author.split(',')[0].strip() if personal else author
    return surname, m.group(2)


ref_keys = [key_of(r) for r in refs]


def cited(surname, year):
    """The first author's surname appears within 80 characters before the year."""
    for m in re.finditer(re.escape(surname), text_part):
        if year in text_part[m.end():m.end() + 80]:
            return True
    return False


uncited = [r for r, k in zip(refs, ref_keys) if not cited(*k)]
check('every reference is cited in the text', not uncited, '; '.join(u[:50] for u in uncited))

missing = []
for m in re.finditer(r'(\d{4}|n\.d\.)(?:-[ab])?\)', text_part):
    year = m.group(1)
    window = text_part[max(0, m.start() - 90):m.start()]
    if '(' in window:
        window = window[window.rfind('('):] if not re.search(r'[A-Z][a-z]+ (?:and|&) [A-Z]', window[window.rfind('('):]) else window
    names = set(re.findall(r"[A-Z][A-Za-z\-’'.]+(?: [A-Z][a-z]+)?", window)) | set(re.findall(r"[A-Z][A-Za-z\-’'.]+", window))
    if not names:
        continue
    if not any(k[0] in names or k[0] in window for k in ref_keys if k[1] == year):
        missing.append(window[-60:].replace('\n', ' ') + year)
missing = [x for x in missing if not re.search(r'(June|October|September|August|May|Act|Version|Surah)\s*\d*\s*$', x[:-4].strip())]
check('every in-text citation is in the reference list', not missing, ' | '.join(missing[:8]))
norm = [re.sub(r'[^a-z]', '', r.lower()) for r in refs]
check('references in alphabetical order (letter by letter)', norm == sorted(norm),
      str([refs[i][:30] for i in range(len(norm) - 1) if norm[i] > norm[i + 1]]))

# ---- stale facts
for pat in (r'\b40 tables', r'84 policies', r'27 migrations', r'\b74 unit', r'\b56 functional', r'Appendix A\.1', r'12 Edge Functions'):
    hits = re.findall(pat, full)
    check(f'no outdated "{pat}"', not hits, str(len(hits)))

# ---- PDF
pdf = pymupdf.open(PDF)
blank = [i + 1 for i, p in enumerate(pdf) if len(p.get_text().strip()) < 8 and not p.get_images()]
check('no blank pages', not blank, str(blank))
fonts = sorted({f[3] for p in pdf for f in p.get_fonts()})
check('fonts embedded in the PDF', all('+' in f for f in fonts), ', '.join(fonts)[:160])
land = [i + 1 for i, p in enumerate(pdf) if p.rect.width > p.rect.height]
check('all pages portrait', not land, str(land))
check('page count', True, str(len(pdf)))
json.dump(results, open('qa_results.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f"\n{sum(1 for r in results if r[1])}/{len(results)} checks passed")
